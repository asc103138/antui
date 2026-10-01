import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D5DD", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def apply_font(run, font_name="微軟正黑體", size_pt=11, bold=False, italic=False, color_rgb=(51,51,51)):
    run.font.name = font_name
    run._r.get_or_add_rPr().set(qn('w:rFonts'), f'{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}hint="eastAsia"')
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is not None:
        rFonts.set(qn('w:eastAsia'), font_name)
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    apply_font(run, font_name="微軟正黑體", size_pt=15, bold=True, color_rgb=(24, 76, 120))
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    apply_font(run, font_name="微軟正黑體", size_pt=12.5, bold=True, color_rgb=(41, 102, 153))
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        apply_font(r_pre, font_name="微軟正黑體", size_pt=10.5, bold=True, color_rgb=(30, 41, 59))
    run = p.add_run(text)
    apply_font(run, font_name="微軟正黑體", size_pt=10.5, bold=False, color_rgb=(51, 65, 85))
    return p

def build_docx(output_path):
    doc = Document()
    
    # 邊界設定
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 頁首大標題
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("臺中市梧棲區中正國小 115 學年度校內語文競賽")
    apply_font(run_title, font_name="微軟正黑體", size_pt=18, bold=True, color_rgb=(15, 44, 89))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("各競賽項目培訓課表、自主練習指南與回家功課手冊（含家長簽章聯）")
    apply_font(run_sub, font_name="微軟正黑體", size_pt=12, bold=False, color_rgb=(71, 85, 105))

    # 壹、時程總覽
    add_heading_1(doc, "壹、競賽基本時程與項目總覽")
    
    schedule_data = [
        ["梯次", "比賽項目", "比賽日期與時間", "時限", "地點", "參賽對象", "題目與評分核心"],
        ["第一梯次", "作文", "115.11.05（四）08:00-09:30", "90分", "視聽教室", "四、五年級", "現場出題，內容50%、結構修辭40%、字體標點10%"],
        ["第一梯次", "寫字（書法）", "115.11.05（四）08:20-09:20", "60分", "圖書室", "四、五年級", "28字指定詩句＋落款，筆法50%、結構章法50%"],
        ["第一梯次", "國語字音字形", "115.11.05（四）12:40-12:50", "10分", "視聽教室", "四、五年級", "200字（音100/形100），112~114歷屆題庫，塗改不計分"],
        ["第二梯次", "臺灣台語朗讀", "115.11.12（四）08:10-09:20", "2分", "校史室", "四、五年級", "四小：阿公變魔術；五小：二抽一。語音45%、氣勢45%、儀態10%"],
        ["第二梯次", "國語朗讀", "115.11.12（四）08:10-09:20", "2分", "簡報室", "四、五年級", "四小：藍蝴蝶；五小：如果我有三天的光明。語音45%/氣勢45%"],
        ["第三梯次", "臺灣台語情境演說", "115.11.19（四）08:10-09:20", "2分+問答", "校史室", "四、五年級", "題目：食中晝。內容、表達、生動、自信、評判即席問答"],
        ["第三梯次", "國語演說", "115.11.19（四）08:10-09:20", "2分", "簡報室", "四、五年級", "二擇一：《同學做了不對的事》《最喜歡的卡通人物》"]
    ]

    table1 = doc.add_table(rows=len(schedule_data), cols=7)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table1)

    col_widths = [Inches(0.8), Inches(1.1), Inches(1.5), Inches(0.6), Inches(0.8), Inches(0.9), Inches(1.6)]

    for row_idx, row_data in enumerate(schedule_data):
        row = table1.rows[row_idx]
        for col_idx, cell_text in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 3, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(cell_text)
            if row_idx == 0:
                set_cell_background(cell, "1E3A8A")
                apply_font(run, font_name="微軟正黑體", size_pt=9.5, bold=True, color_rgb=(255, 255, 255))
            else:
                if row_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                apply_font(run, font_name="微軟正黑體", size_pt=8.5, bold=(col_idx == 1), color_rgb=(30, 41, 59))

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 貳、在校微型集訓課表
    add_heading_1(doc, "貳、在校「微型集訓課表」每週輪替表")
    p_desc = doc.add_paragraph()
    r = p_desc.add_run("採「在校微型集訓（每次 25~35 分鐘）＋ 每日課後精準練習（15~20 分鐘）」雙軌制，不佔用常規學科課程。")
    apply_font(r, font_name="微軟正黑體", size_pt=10, italic=True, color_rgb=(100, 116, 139))

    timetable_data = [
        ["集訓時段", "星期一", "星期二", "星期三", "星期四", "星期五"],
        [
            "早自修\n08:00-08:35\n(35分鐘)",
            "【國語朗讀】\n正音、斷句標記\n【字音字形】\n限時快閃小測驗",
            "【國語演說】\n講稿骨架與背誦\n【寫字組】\n單字結構分析",
            "【全體筆試模擬】\n作文審題構思\n字音字形題庫刷題",
            "【台語朗讀】\n台羅拼音與連讀變調\n【寫字組】\n全張排版折格指導",
            "【台語情境演說】\n口語表達與生動身段\n即席問答應對指導"
        ],
        [
            "午休前段\n12:40-13:15\n(35分鐘)",
            "【寫字組】\n毛筆運筆、提按精練\n九宮格重點字精雕",
            "【台語情境演說】\n情境動作與語調雕琢\n評判現場抽問演練",
            "（全校教師研習）\n選手自主背誦自學\n不排定集中培訓",
            "【字音字形衝刺】\n10分鐘全真模擬卷\n錯字本即時檢核",
            "【朗讀演說驗收】\n模擬登台走位禮儀\n2分鐘鈴聲下台訓練"
        ]
    ]

    table2 = doc.add_table(rows=len(timetable_data), cols=6)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table2)
    tt_widths = [Inches(1.2), Inches(1.22), Inches(1.22), Inches(1.22), Inches(1.22), Inches(1.22)]

    for row_idx, row_data in enumerate(timetable_data):
        row = table2.rows[row_idx]
        for col_idx, cell_text in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.width = tt_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 or row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(cell_text)
            if row_idx == 0:
                set_cell_background(cell, "2563EB")
                apply_font(run, font_name="微軟正黑體", size_pt=9.5, bold=True, color_rgb=(255, 255, 255))
            else:
                if col_idx == 0:
                    set_cell_background(cell, "EFF6FF")
                    apply_font(run, font_name="微軟正黑體", size_pt=9, bold=True, color_rgb=(30, 58, 138))
                else:
                    apply_font(run, font_name="微軟正黑體", size_pt=8.5, bold=False, color_rgb=(30, 41, 59))

    doc.add_page_break()

    # 參、各競賽項目專屬訓練進度與回家功課
    add_heading_1(doc, "參、各競賽項目培訓進度、練習項目與回家功課")

    # 1. 寫字
    add_heading_2(doc, "【項目一：寫字（書法）組】")
    add_bullet(doc, "比賽時間：115.11.05（四）08:20-09:20（60 分鐘，圖書室）。", "［競賽規格］")
    add_bullet(doc, "四尺手工宣紙四開（70cm × 35cm）一張，傳統毛筆楷書書寫，錯漏字每字扣 3 分，未寫完扣 2 分。", "［規則評分］")
    add_bullet(doc, "「拒絕毒品愛自己 反對霸凌不旁觀 減速慢行少意外 共創和諧好未來」＋ 落款（中正國小○年○班 ○○○書）。", "［指定題目］")
    add_bullet(doc, "第1週字形重心（霸、毒、凌、諧）→ 第2週提按筆法 → 第3週折格章法 → 第4週落款協調 → 第5週60分全真模擬。", "［分週進度］")
    add_bullet(doc, "在校練習：單字九宮格精雕、折格技巧（四行每行7字＋落款）、墨色濃淡控制。", "［在校練習］")
    add_bullet(doc, "週一至週四精練指定句 7 字（每字 3 遍）；週末完成全張宣紙模擬 1 次送交老師批閱。", "［回家功課］")

    # 2. 作文
    add_heading_2(doc, "【項目二：作文組】")
    add_bullet(doc, "比賽時間：115.11.05（四）08:00-09:30（90 分鐘，視聽教室）。", "［競賽規格］")
    add_bullet(doc, "題目現場公佈，學校發給稿紙，禁詩歌韻文、禁鉛筆紅筆。評分：內容思想50%、結構修辭40%、字體標點10%。", "［規則評分］")
    add_bullet(doc, "第1週審題題眼 → 第2週起承轉合四段骨架 → 第3週五感摹寫與修辭 → 第4週配速（5分大綱/75分行文/10分檢查）→ 第5週全真衝刺。", "［分週進度］")
    add_bullet(doc, "在校練習：3分鐘心智圖大綱速成、亮點段落改寫手術（動詞＋譬喻＋情意昇華）。", "［在校練習］")
    add_bullet(doc, "週一至週四每日摘錄 2 個成語＋1 句優美句子並造段；週末計時 75 分鐘完成一篇 600~800 字命題作文。", "［回家功課］")

    # 3. 國語字音字形
    add_heading_2(doc, "【項目三：國語字音字形組】")
    add_bullet(doc, "比賽時間：115.11.05（四）12:40-12:50（10 分鐘，視聽教室）。", "［競賽規格］")
    add_bullet(doc, "200 字（字音 100、字形 100），全數出自 112、113、114 年全國語文競賽試卷。每字 0.5 分，塗改一律不計分！", "［規則評分］")
    add_bullet(doc, "第1週112年試卷精熟 → 第2週113年題庫攻堅 → 第3週114年考題剖析 → 第4週跨年混合盲測 → 第5週10分鐘極限配速。", "［分週進度］")
    add_bullet(doc, "在校練習：5分鐘快閃抽測 50 題、標準字體陷阱拆解（如橫折、瞥、肺、掣）。", "［在校練習］")
    add_bullet(doc, "週一至週四每日手寫訂正 20 題（10音+10形），家長口頭抽考 5 題；週末進行 100 題（5分鐘）全真自測。", "［回家功課］")

    # 4. 國語朗讀
    add_heading_2(doc, "【項目四：國語朗讀組】")
    add_bullet(doc, "比賽時間：115.11.12（四）08:10-09:20（限時 2 分鐘，活動中心簡報室）。", "［競賽規格］")
    add_bullet(doc, "四年級：林清玄〈藍蝴蝶〉；五年級：海倫凱勒〈如果我有三天的光明〉。語音45%、氣勢45%、儀態10%。", "［指定篇目］")
    add_bullet(doc, "第1週正音斷句符號 → 第2週聲情層次與文氣 → 第3週咬字共鳴與丹田發聲 → 第4週2分鐘配速（約400字）→ 第5-6週上下台禮儀。", "［分週進度］")
    add_bullet(doc, "在校練習：調值聽辨（一聲高平、四聲全降）、捧稿 45 度角與眼神巡視評審。", "［在校練習］")
    add_bullet(doc, "週一至週四每日站姿大聲朗讀 3 次，手機錄音 1 次自檢；週末在家長面前朗讀 2 分鐘，家長簽核評分。", "［回家功課］")

    # 5. 臺灣台語朗讀
    add_heading_2(doc, "【項目五：臺灣台語朗讀組】")
    add_bullet(doc, "比賽時間：115.11.12（四）08:10-09:20（限時 2 分鐘，活動中心校史室）。", "［競賽規格］")
    add_bullet(doc, "四年級：〈04 阿公變魔術〉；五年級：二抽一（〈02 膨鼠著生驚〉、〈04 阿公變魔術〉）。語音45%/氣勢45%/儀態10%。", "［指定篇目］")
    add_bullet(doc, "第1週官方音檔Shadowing跟讀 → 第2週台羅拼音與連讀變調 → 第3週口氣韻味（祖孫情/生驚動態）→ 第4週二抽一盲抽 → 第5-6週控速。", "［分週進度］")
    add_bullet(doc, "在校練習：台語長句變調糾錯、入聲字（-p, -t, -k, -h）俐落收音、斷詞呼吸法。", "［在校練習］")
    add_bullet(doc, "週一至週四跟讀官方 MP3 音檔 2 次＋自主朗讀 2 次；週末錄製 2 分鐘音檔予長輩或老師聆聽並簽名。", "［回家功課］")

    # 6. 國語演說
    add_heading_2(doc, "【項目六：國語演說組】")
    add_bullet(doc, "比賽時間：115.11.19（四）08:10-09:20（限時 2 分鐘，活動中心簡報室）。", "［競賽規格］")
    add_bullet(doc, "題目二擇一：《如果同學做了不對的事》、《我最喜歡的卡通人物》。1分30秒一鈴、2分二鈴停止。內容50%/語音40%/儀態10%。", "［指定題目］")
    add_bullet(doc, "第1週黃金420字講稿定稿 → 第2週分段骨架記憶脫稿 → 第3週聲情輕重起伏 → 第4週自然手勢設計 → 第5-7週1分45秒配速與抗壓。", "［分週進度］")
    add_bullet(doc, "在校練習：分段抽背、鈴聲聽聞反應（1:30一響進結論、2:00二響俐落鞠躬致謝下台）。", "［在校練習］")
    add_bullet(doc, "週一至週四鏡前脫稿演練 3 遍，控時 1 分 45 秒至 1 分 55 秒；週末請家長錄影回放檢核表情與口頭禪。", "［回家功課］")

    # 7. 臺灣台語情境演說
    add_heading_2(doc, "【項目七：臺灣台語情境式演說組】")
    add_bullet(doc, "比賽時間：115.11.19（四）08:10-09:20（限時 2 分鐘＋問答，活動中心校史室）。", "［競賽規格］")
    add_bullet(doc, "題目：〈食中晝〉（阿明排隊添飯無聊耍弄、菜桶偃倒、提布鑢仔鑢清氣＋評判委員現場即席台語問答）。", "［情境題目］")
    add_bullet(doc, "第1週台詞定稿融入俚語 → 第2週生動身段動作 → 第3週脫稿流暢度 → 第4-5週Q&A題庫抽問集訓 → 第6-7週全真模擬。", "［分週進度］")
    add_bullet(doc, "在校練習：四格看圖快說、評判委員模擬抽問（應急處理、反省、同學互助）。", "［在校練習］")
    add_bullet(doc, "週一至週四向家人用台語講述故事 2 遍並回答日常台語問題；週末錄製「2分鐘演說＋1題問答」影音供家長簽審。", "［回家功課］")

    doc.add_page_break()

    # 肆、賽前應變與檢核清單
    add_heading_1(doc, "肆、選手臨場應變心法與賽前裝備檢核清單")
    
    add_bullet(doc, "忘詞自救法：切忌吐舌、抓頭或嘆氣！保持微笑注視評審，深吸一口氣，直接跳接下一個記得的關鍵詞，自然接續。", "［心理素質］")
    add_bullet(doc, "鈴聲掌握：聽聞一響鈴（剩30秒）立即收束進入結論；聞二響鈴（時間到）講完該句立刻鞠躬「謝謝評判老師」退場。", "［時間掌控］")
    add_bullet(doc, "字音字形零塗改：超過 3 秒立刻跳題！嚴禁塗改或描字，確保卷面整潔與零扣分。", "［作答紀律］")

    p_chk = doc.add_paragraph()
    p_chk.paragraph_format.space_before = Pt(8)
    p_chk.paragraph_format.space_after = Pt(4)
    r_chk = p_chk.add_run("【各組比賽當日前一晚必備物品打勾檢核表】")
    apply_font(r_chk, font_name="微軟正黑體", size_pt=11, bold=True, color_rgb=(30, 58, 138))

    chk_data = [
        ["組別", "選手自備物品與裝備清單", "賽前一晚確認"],
        ["寫字（書法）組", "兼毫/狼毫大楷筆（洗淨開鋒）、黑色厚毛氈墊布、墨汁、墨碟、鎮尺、吸水布、落款小楷筆", "［  ］已備妥"],
        ["作文組", "深藍色/黑色原子筆（0.5mm~0.7mm 準備 2 支）、透明墊板、水壺（嚴禁鉛筆紅筆）", "［  ］已備妥"],
        ["國語字音字形組", "深藍/黑色滑順原子筆 2 支、透明筆袋（嚴禁使用修正帶、立可白）", "［  ］已備妥"],
        ["國語/台語朗讀組", "自備標記註音朗讀稿（候場溫習用）、水壺、乾淨整齊學校制服/運動服", "［  ］已備妥"],
        ["國語/台語演說組", "演說講稿卡（賽前候場默背）、水壺、儀容整潔（頭髮不遮眼、鞋帶繫緊）", "［  ］已備妥"]
    ]

    table3 = doc.add_table(rows=len(chk_data), cols=3)
    table3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table3)
    chk_widths = [Inches(1.5), Inches(4.5), Inches(1.2)]

    for row_idx, row_data in enumerate(chk_data):
        row = table3.rows[row_idx]
        for col_idx, cell_text in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.width = chk_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(cell_text)
            if row_idx == 0:
                set_cell_background(cell, "0F766E")
                apply_font(run, font_name="微軟正黑體", size_pt=9.5, bold=True, color_rgb=(255, 255, 255))
            else:
                if row_idx % 2 == 1:
                    set_cell_background(cell, "F0FDFA")
                apply_font(run, font_name="微軟正黑體", size_pt=9, bold=(col_idx==0), color_rgb=(30, 41, 59))

    doc.add_page_break()

    # 伍、家長簽章聯與每週進度檢核表
    add_heading_1(doc, "伍、選手自主練習進度與家長每週簽章回饋聯")
    
    p_letter = doc.add_paragraph()
    r_letter = p_letter.add_run("親愛的家長您好：恭喜貴子弟獲選為本校 115 學年度校內語文競賽選手！語文能力的深化來自「每日少量恆毅的堅持」。學校老師已利用在校零碎時間進行精準指導，懇請您每日撥冗 15 分鐘陪伴、聆聽或抽考，並於下方回饋聯簽核，親師攜手為孩子打造自信發光的舞台！")
    apply_font(r_letter, font_name="微軟正黑體", size_pt=9.5, italic=False, color_rgb=(71, 85, 105))

    # 產生兩週格式的家長簽聯（可剪裁列印）
    for w in [1, 2]:
        p_week = doc.add_paragraph()
        p_week.paragraph_format.space_before = Pt(10)
        p_week.paragraph_format.space_after = Pt(4)
        rw = p_week.add_run(f"【第 {w} 週選手自主練習與家庭檢核聯】　班級：____年____班　姓名：__________　項目：______________")
        apply_font(rw, font_name="微軟正黑體", size_pt=10.5, bold=True, color_rgb=(15, 23, 42))

        eval_data = [
            ["日 期", "每日在校與回家練習重點", "完成時間", "學生自評", "家長簽章確認"],
            ["星期一", "完成當日項目指定練習（如單字精雕 / 朗讀3遍 / 20題字音字形）", "___ 分鐘", "［ ］精熟 ［ ］尚可", "簽章：__________"],
            ["星期二", "完成當日項目指定練習（如講稿背誦 / 成語造句 / 變調朗讀）", "___ 分鐘", "［ ］精熟 ［ ］尚可", "簽章：__________"],
            ["星期三", "自主加強或整理錯題筆記 / 錄音回聽自我診斷", "___ 分鐘", "［ ］精熟 ［ ］尚可", "簽章：__________"],
            ["星期四", "完成當日項目指定練習（如全篇折格 / 講稿手勢 / 抽測訂正）", "___ 分鐘", "［ ］精熟 ［ ］尚可", "簽章：__________"],
            ["週末驗收", "完成 1 次全真模擬（限時寫字 / 75分作文 / 2分脫稿演說 / 錄音）", "___ 分鐘", "［ ］精熟 ［ ］尚可", "簽章：__________"]
        ]

        table_w = doc.add_table(rows=len(eval_data), cols=5)
        table_w.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table_w)
        w_widths = [Inches(1.0), Inches(3.2), Inches(1.0), Inches(1.0), Inches(1.0)]

        for row_idx, row_data in enumerate(eval_data):
            row = table_w.rows[row_idx]
            for col_idx, cell_text in enumerate(row_data):
                cell = row.cells[col_idx]
                cell.width = w_widths[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
                
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 2, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                
                run = p.add_run(cell_text)
                if row_idx == 0:
                    set_cell_background(cell, "334155")
                    apply_font(run, font_name="微軟正黑體", size_pt=9, bold=True, color_rgb=(255, 255, 255))
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
                    apply_font(run, font_name="微軟正黑體", size_pt=8.5, bold=False, color_rgb=(30, 41, 59))

        p_feedback = doc.add_paragraph()
        p_feedback.paragraph_format.space_before = Pt(4)
        p_feedback.paragraph_format.space_after = Pt(12)
        rf1 = p_feedback.add_run("★ 家長溫馨回饋／鼓勵小語：________________________________________________________\n★ 指導老師評語／訓練叮嚀：________________________________________________________")
        apply_font(rf1, font_name="微軟正黑體", size_pt=8.5, italic=True, color_rgb=(100, 116, 139))

    doc.save(output_path)
    print("Document successfully created at:", output_path)

if __name__ == "__main__":
    out = os.path.abspath("115學年度校內語文競賽_各項目選手訓練課表與回家功課手冊(含家長簽章聯).docx")
    build_docx(out)
