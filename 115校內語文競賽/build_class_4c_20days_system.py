# -*- coding: utf-8 -*-
"""
臺中市梧棲區中正國小 115 學年度校內語文競賽
【四年丙班】專屬 7 大項目 10/5 ~ 10/30 (Day 01 至 Day 20)
老師指導指引與評量標準 ＋ 學生挖空練習學習單與自評表 生成系統
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.abspath("四年丙班_各項目培訓資料庫")
os.makedirs(BASE_DIR, exist_ok=True)

# 20 天日期對照
DATES_20 = [
    ("Day 01", "10/05 (一)"),
    ("Day 02", "10/06 (二)"),
    ("Day 03", "10/07 (三)"),
    ("Day 04", "10/08 (四)"),
    ("Day 05", "10/09 (五)"),
    ("Day 06", "10/12 (一)"),
    ("Day 07", "10/13 (二)"),
    ("Day 08", "10/14 (三)"),
    ("Day 09", "10/15 (四)"),
    ("Day 10", "10/16 (五)"),
    ("Day 11", "10/19 (一)"),
    ("Day 12", "10/20 (二)"),
    ("Day 13", "10/21 (三)"),
    ("Day 14", "10/22 (四)"),
    ("Day 15", "10/23 (五)"),
    ("Day 16", "10/26 (一)"),
    ("Day 17", "10/27 (二)"),
    ("Day 18", "10/28 (三)"),
    ("Day 19", "10/29 (四)"),
    ("Day 20", "10/30 (五)")
]

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D5DD", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def apply_font(run, font_name="標楷體", size_pt=16, bold=False, italic=False, color_rgb=(51,51,51)):
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    run.font.size = Pt(16)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def create_doc():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '標楷體'
    style.font.size = Pt(16)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '標楷體')
    rFonts.set(qn('w:ascii'), '標楷體')
    rFonts.set(qn('w:hAnsi'), '標楷體')
    for s in doc.sections:
        s.top_margin = Inches(0.65)
        s.bottom_margin = Inches(0.65)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)
    return doc

def add_header(doc, title, subtitle):
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run(title)
    apply_font(r1, size_pt=16, bold=True, color_rgb=(15, 44, 89))

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(10)
    r2 = p2.add_run(subtitle)
    apply_font(r2, size_pt=16, bold=True, color_rgb=(194, 65, 12))

def add_section_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    apply_font(r, size_pt=16, bold=True, color_rgb=(15, 44, 89))
    return p

def add_section_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    apply_font(r, size_pt=16, bold=True, color_rgb=(30, 64, 175))
    return p

def add_bullet_pt(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        apply_font(r_pre, size_pt=16, bold=True, color_rgb=(30, 41, 59))
    r = p.add_run(text)
    apply_font(r, size_pt=16, bold=False, color_rgb=(51, 65, 85))
    return p

def add_simple_table(doc, data, widths, header_color="1E3A8A"):
    t = doc.add_table(rows=len(data), cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t)
    for r_i, row_data in enumerate(data):
        row = t.rows[r_i]
        for c_i, cell_text in enumerate(row_data):
            cell = row.cells[c_i]
            cell.width = widths[c_i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=50, bottom=50, left=50, right=50)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (r_i==0 or c_i==0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(cell_text)
            if r_i == 0:
                set_cell_background(cell, header_color)
                apply_font(run, size_pt=16, bold=True, color_rgb=(255, 255, 255))
            else:
                if r_i % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                apply_font(run, size_pt=16, bold=(c_i==0), color_rgb=(30, 41, 59))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

print("Base setup loaded.")
